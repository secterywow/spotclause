import os
import tempfile
from celery import shared_task
from app.celery_app import celery_app
from app.logger import get_logger

logger = get_logger(__name__)


@celery_app.task(bind=True, max_retries=3)
def parse_file(self, file_path: str, mime_type: str):
    """Parse a file and extract text content."""
    try:
        logger.info("Parsing file", file_path=file_path, mime_type=mime_type)

        if mime_type == 'application/pdf':
            return parse_pdf(file_path)
        elif mime_type in ['application/msword', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document']:
            return parse_word(file_path)
        elif mime_type.startswith('image/'):
            return parse_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {mime_type}")

    except Exception as exc:
        logger.error("File parsing failed", error=str(exc), file_path=file_path)
        self.retry(exc=exc, countdown=5)
    finally:
        # Clean up temp file
        if os.path.exists(file_path):
            os.remove(file_path)


def parse_pdf(file_path: str) -> dict:
    """Extract text from PDF using pdfplumber."""
    import pdfplumber

    text = ""
    page_count = 0

    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n\n"

    return {
        "text": text.strip(),
        "page_count": page_count,
        "type": "pdf",
    }


def parse_word(file_path: str) -> dict:
    """Extract text from Word document using python-docx."""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    text = "\n\n".join(paragraphs)

    return {
        "text": text.strip(),
        "page_count": len(doc.sections),
        "type": "word",
    }


def parse_image(file_path: str) -> dict:
    """Extract text from image using pytesseract (Tesseract OCR)."""
    from PIL import Image
    import pytesseract

    image = Image.open(file_path)
    text = pytesseract.image_to_string(image, lang='eng+chi_sim+chi_tra')

    return {
        "text": text.strip(),
        "page_count": 1,
        "type": "image",
    }
